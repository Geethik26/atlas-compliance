"""Privacy-minimizing employee workbook loader."""

from __future__ import annotations

from collections.abc import Iterator
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .models import Employee

ALLOWED_EMPLOYEE_FIELDS = (
    "employee_id",
    "work_country",
    "work_state",
    "work_location_code",
    "pay_basis",
    "hourly_rate_ast",
    "annual_salary_ast",
    "scheduled_hours_per_week",
    "currency",
    "employment_status",
    "minimum_wage_coverage",
)


def _text(value: Any) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _decimal(value: Any) -> Decimal | None:
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    try:
        return Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise ValueError(f"Invalid numeric employee value: {value!r}") from exc


def load_employees(path: Path) -> list[Employee]:
    """Load only allowlisted compliance fields from an XLSX workbook.

    Sensitive and unrelated columns are never copied into row dictionaries or
    domain objects.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows: Iterator[tuple[Any, ...]] = worksheet.iter_rows(values_only=True)
        try:
            headers = next(rows)
        except StopIteration as exc:
            raise ValueError("Employee workbook is empty") from exc

        positions = {str(name): index for index, name in enumerate(headers)}
        missing = [name for name in ALLOWED_EMPLOYEE_FIELDS if name not in positions]
        if missing:
            raise ValueError(f"Employee workbook is missing columns: {missing}")

        employees: list[Employee] = []
        for row in rows:
            selected = {
                name: row[positions[name]] if positions[name] < len(row) else None
                for name in ALLOWED_EMPLOYEE_FIELDS
            }
            employees.append(
                Employee(
                    employee_id=_text(selected["employee_id"]),
                    work_country=_text(selected["work_country"]),
                    work_state=_text(selected["work_state"]),
                    work_location_code=_text(selected["work_location_code"]),
                    pay_basis=_text(selected["pay_basis"]),
                    hourly_rate_ast=_decimal(selected["hourly_rate_ast"]),
                    annual_salary_ast=_decimal(selected["annual_salary_ast"]),
                    scheduled_hours_per_week=_decimal(
                        selected["scheduled_hours_per_week"]
                    ),
                    currency=_text(selected["currency"]),
                    employment_status=_text(selected["employment_status"]),
                    minimum_wage_coverage=_text(
                        selected["minimum_wage_coverage"]
                    ),
                )
            )
        return employees
    finally:
        workbook.close()
